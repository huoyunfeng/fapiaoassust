"""端到端验证：用户「餐费报销明细表」两种导入方式。"""
import io
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, ".")
from openpyxl import Workbook
from openpyxl.styles import Alignment


def make_user_excel():
    """复现用户 Excel：第1行合并标题、第2行表头、第3行起数据。"""
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:H1")
    ws["A1"] = "餐费报销明细表"
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["序号", "部门", "姓名", "实际用餐日期", "星期", "项目名称", "金额", "备注"])
    ws.append([1, "工程二事业部", "霍云峰", "2026/3/23", "星期一",
               "软件产品及服务采购合同", 31.8, ""])
    ws.append([2, "工程二事业部", "霍云峰", "2026/3/24", "星期二",
               "软件产品及服务采购合同", 45.0, "客户招待"])
    ws.append([3, "工程一部", "李四", "2026/3/25", "星期三",
               "办公用品采购", 88.5, ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    content = make_user_excel()
    tmp = tempfile.mkdtemp()
    from app.config import settings
    settings.db_path = Path(tmp) / "t.db"
    from app.repository.db import init_schema
    init_schema()

    from app.infra.excel import read_headers
    from app.service.mapping_service import mapping_service
    from app.service import reimburse_service

    print("=" * 60)
    print("A. 自定义映射导入（模拟前端向导）")
    print("=" * 60)
    headers, hrow = read_headers(content)
    print("表头:", headers, "| 表头行:", hrow)
    det = mapping_service.detect(headers, header_row=hrow)
    mp = {h: (det["mapping"].get(h) or "custom_" + h) for h in headers}
    print("向导映射:", json.dumps(mp, ensure_ascii=False))
    fc = mapping_service.build_fields_config(headers, mp, hrow)
    print("custom 字段:", [f["excel_header"] for f in fc["fields"] if f["is_custom"]])
    res = reimburse_service.import_excel("餐费报销明细表.xlsx", content,
                                         fields_config=fc, remember_mapping=True)
    print("结果:", json.dumps(res, ensure_ascii=False, default=str))
    for r in reimburse_service.list_reimburses():
        print("   #%s %s 金额=%s 费用=%s 日期=%s 附加=%s" % (
            r["id"], r["reimburse_no"], r["amount"], r["expense_type"],
            r["reimburse_date"], json.dumps(r["extra_data"], ensure_ascii=False)))

    print()
    print("=" * 60)
    print("B. 二次导入应命中映射缓存")
    print("=" * 60)
    h2, hr2 = read_headers(content)
    d2 = mapping_service.detect(h2, header_row=hr2)
    print("命中缓存:", d2["from_cache"], "| 模板 id:", d2["template_id"])

    print()
    print("=" * 60)
    print("C. 模板模式导入（经典模板）")
    print("=" * 60)
    tpls = reimburse_service.template_service.list() if hasattr(reimburse_service, "template_service") else []
    from app.service.template_service import template_service
    classic = [t for t in template_service.list() if t["is_system"]][0]
    res3 = reimburse_service.import_excel("餐费报销明细表.xlsx", content, template_id=classic["id"])
    print("经典模板结果:", json.dumps(res3, ensure_ascii=False, default=str))

    print()
    print("=" * 60)
    print("D. 手动录入")
    print("=" * 60)
    print(reimburse_service.manual_create(
        {"reimburse_no": "BX-001", "submitter": "王五", "amount": "120.5",
         "expense_type": "差旅费", "reimburse_date": "2026-09-01"}))


if __name__ == "__main__":
    main()
